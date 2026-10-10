"""API Admin 1.3 – Data Preview và Export Excel/PDF."""
from __future__ import annotations

import time
from io import BytesIO

from openpyxl import load_workbook

from backend.app.core.config import get_settings
from tests.backend.conftest import API

PREVIEW = f"{API}/admin/reports/preview"
EXPORT = f"{API}/admin/reports/export"


def test_preview_enrollment(client, admin_headers, courses):
    body = client.get(PREVIEW, params={"report_type": "enrollment", "page_size": 200}, headers=admin_headers).json()
    assert body["total"] == len(courses) == len(body["items"])
    assert body["filters"]["academic_year"] == "2026-2027" and body["filters"]["semester"] == 1
    keys = [c["key"] for c in body["columns"]]
    assert keys[:4] == ["faculty_name", "category_name", "course_shortname", "course_fullname"]
    assert set(keys) <= set(body["items"][0])
    assert sum(i["total_enrolled_students"] for i in body["items"]) == sum(c["enrolled"] for c in courses)

    one = client.get(PREVIEW, params={"report_type": "enrollment", "course_id": courses[0]["course_id"]},
                     headers=admin_headers).json()
    assert one["total"] == 1
    none = client.get(PREVIEW, params={"report_type": "enrollment", "faculty_id": 999_999}, headers=admin_headers).json()
    assert none["total"] == 0 and none["items"] == []


def test_preview_risk(client, admin_headers, courses):
    everyone = client.get(PREVIEW, params={"report_type": "risk", "page_size": 10}, headers=admin_headers).json()
    assert everyone["total"] == sum(c["enrolled"] for c in courses)
    assert len(everyone["items"]) == 10
    red = client.get(PREVIEW, params={"report_type": "risk", "risk_level": "RED"}, headers=admin_headers).json()
    assert red["total"] == sum(c["red_count"] for c in courses)
    assert {i["risk_level"] for i in red["items"]} == {"RED"}


def test_preview_traffic_and_validation(client, admin_headers):
    body = client.get(PREVIEW, params={"report_type": "traffic", "from": "2026-10-01", "to": "2026-10-02"},
                      headers=admin_headers).json()
    assert body["total"] == 48 and body["filters"] == {"from": "2026-10-01", "to": "2026-10-02"}
    assert client.get(PREVIEW, headers=admin_headers).status_code == 422
    assert client.get(PREVIEW, params={"report_type": "grades"}, headers=admin_headers).status_code == 422
    assert client.get(PREVIEW, params={"report_type": "traffic", "from": "2025-01-01", "to": "2026-10-01"},
                      headers=admin_headers).status_code == 422
    assert client.get(PREVIEW, params={"report_type": "risk", "page_size": 201}, headers=admin_headers).status_code == 422


def test_export_enrollment_xlsx(client, admin_headers, courses):
    res = client.get(EXPORT, params={"report_type": "enrollment", "format": "xlsx"}, headers=admin_headers)
    assert res.status_code == 200
    assert 'filename="bao_cao_enrollment_HK1_2026-2027.xlsx"' in res.headers["content-disposition"]
    assert res.headers["x-data-updated-at"]
    workbook = load_workbook(BytesIO(res.content))
    rows = list(workbook["Báo cáo"].iter_rows(values_only=True))
    assert rows[0][:2] == ("Khoa", "Bộ môn")
    assert len(rows) - 1 == len(courses)
    assert isinstance(rows[1][4], int) and isinstance(rows[1][7], float)          # giữ kiểu số trong Excel
    info = {r[0]: r[1] for r in workbook["Thông tin"].iter_rows(values_only=True)}
    assert info["Năm học – học kỳ"] == "2026-2027 – HK1" and info["Số dòng"] == len(courses)


def test_export_risk_pdf(client, admin_headers, courses):
    # Toàn trường (5 khoa, > 5.000 dòng) vượt giới hạn PDF -> 422 kèm hướng dẫn thu hẹp bộ lọc / xuất Excel
    whole = client.get(EXPORT, params={"report_type": "risk", "format": "pdf"}, headers=admin_headers)
    assert whole.status_code == 422 and "xuất Excel" in whole.json()["detail"]
    # Khoa lớn nhất nằm trong giới hạn -> xuất được
    enrolled: dict[int, int] = {}
    for c in courses:
        enrolled[c["faculty_id"]] = enrolled.get(c["faculty_id"], 0) + c["enrolled"]
    faculty_id, rows = max(enrolled.items(), key=lambda x: x[1])
    started = time.perf_counter()
    res = client.get(EXPORT, params={"report_type": "risk", "format": "pdf", "faculty_id": faculty_id},
                     headers=admin_headers)
    elapsed = time.perf_counter() - started
    assert res.status_code == 200, res.text
    assert res.headers["content-type"] == "application/pdf" and res.content.startswith(b"%PDF")
    assert len(res.content) > 50_000
    assert elapsed < 30, f"Xuất PDF {rows} dòng mất {elapsed:.1f}s"


def test_export_traffic_xlsx(client, admin_headers):
    res = client.get(EXPORT, params={"report_type": "traffic", "from": "2026-10-01", "to": "2026-10-07"},
                     headers=admin_headers)
    rows = list(load_workbook(BytesIO(res.content), read_only=True)["Báo cáo"].iter_rows(values_only=True))
    assert len(rows) - 1 == 7 * 24
    assert 'bao_cao_traffic_20261001_20261007.xlsx' in res.headers["content-disposition"]


def test_export_row_limit(client, admin_headers, monkeypatch):
    monkeypatch.setenv("BACKEND_PDF_MAX_ROWS", "10")
    get_settings.cache_clear()
    try:
        res = client.get(EXPORT, params={"report_type": "risk", "format": "pdf"}, headers=admin_headers)
        assert res.status_code == 422 and "xuất Excel" in res.json()["detail"]
        assert client.get(EXPORT, params={"report_type": "enrollment", "format": "pdf"},
                          headers=admin_headers).status_code == 422
    finally:
        monkeypatch.undo()
        get_settings.cache_clear()
